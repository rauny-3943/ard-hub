#!/usr/bin/env python3
"""Build the three deliverables (HTML map, CSV/XLSX reading table, BibLaTeX) from data/corpus.json.

Outputs:
  site/index.html                 stand-alone page published on GitHub Pages
  site/ARD_reading_table.{csv,xlsx}
  site/ARD_corpus.bib
  build/ard_map_fragment.html     same page in Claude artifact format (no <html>/<head> wrapper)
The script exits with an error if the deliverables do not contain exactly the same keys and DOIs.
"""
import collections
import csv
import html
import json
import re
import sys
import unicodedata
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / 'data'
OUT = ROOT / 'site'          # published on GitHub Pages
FRAG = ROOT / 'build'        # fragment version, to republish the map as a Claude artifact
OUT.mkdir(exist_ok=True); FRAG.mkdir(exist_ok=True)
D = json.load(open(DATA / 'corpus.json', encoding='utf-8'))
META, REFS = D['meta'], D['refs']
NOTES = []  # interpretation choices made at build time

MAIN_THEMES = ['Geomorphology & ARD', 'Hydrochemistry & ARD', 'Soil & ARD']
REGIONS = ['To be specified', 'Cordillera Blanca', 'Pyrenees', 'Alps', 'Rocky Mountains', 'Arctic & subarctic', 'Other Andes',
           'Antarctica', 'Himalaya & Tibet', 'Other', 'Global / no location']
LEVELS = ['full text', 'abstract', 'metadata']

# ---------- 1. checks and enrichment ----------
for x in REFS:
    bad = [a for a in x['authors'] if ',' not in a]
    if bad:
        print(f"Warning: authors not in 'Last, First' form in {x['key']}: {bad}", file=sys.stderr)

# reading notes from PDFs (data/notes/<key>.json, written by pdf_notes.py and then reviewed)
NOTE_FIELDS = ['question', 'site', 'methods', 'results', 'contribution', 'reading_level']
n_notes = 0
for f in sorted((DATA / 'notes').glob('*.json')):
    nt = json.loads(f.read_text(encoding='utf-8'))
    x = next((r for r in REFS if r['key'] == nt.get('key', f.stem)), None)
    if x is None:
        print(f"Note skipped (key not in corpus): {f.name}", file=sys.stderr); continue
    for c in NOTE_FIELDS:
        if nt.get(c): x[c] = nt[c]
    msg = (f"Note completed on {nt.get('date', '?')} by {nt.get('generated_by', 'an LLM')} from the full text"
           + (" (reviewed)." if nt.get('reviewed') else " (to be reviewed)."))
    x['note'] = (x.get('note', '') + ' ' + msg).strip()
    n_notes += 1
if n_notes:
    NOTES.append(f"{n_notes} reading note(s) added from data/notes/ (full text, summarised by an LLM).")

# de-duplication: DOI, then normalised title
def norm_title(t):
    t = unicodedata.normalize('NFKD', t.lower())
    return re.sub(r'[^a-z0-9]', '', t)[:80]
seen_doi, seen_t, dups, keep = {}, {}, [], []
for x in REFS:
    d = x['doi'].lower().strip()
    t = norm_title(x['title'])
    if (d and d in seen_doi) or t in seen_t:
        dups.append((x['key'], seen_doi.get(d) or seen_t.get(t))); continue
    if d: seen_doi[d] = x['key']
    seen_t[t] = x['key']; keep.append(x)
REFS = keep
if dups:
    sys.exit(f'Duplicates found (key, duplicate of): {dups}')
KEYS = {x['key'] for x in REFS}
BY = {x['key']: x for x in REFS}
for x in REFS:
    if x['region'] not in REGIONS:
        sys.exit(f"Unknown region for {x['key']}: {x['region']!r}. Use one of: {REGIONS}")

# cited by = inverse of "cites"
for x in REFS: x['cites'] = [k for k in x['cites'] if k in KEYS and k != x['key']]
cited_by = collections.defaultdict(list)
for x in REFS:
    for k in x['cites']: cited_by[k].append(x['key'])
for x in REFS:
    x['cited_by'] = sorted(cited_by[x['key']], key=lambda k: (BY[k]['year'], k))
    x['cites'] = sorted(x['cites'], key=lambda k: (BY[k]['year'], k))

# themes: four families, and ad hoc "Other – …" labels grouped
OTHER_MAP = [
    (r'mechanism|mineralog', 'mechanisms and mineralogy'),
    (r'microbio', 'microbiology'),
    (r'aquatic ecology|ecotox|^limnology$', 'aquatic ecology'),
    (r'plant|soil acidification|pedogenesis', 'vegetation, soils and nutrition (context)'),
    (r'review|teaching', 'review/synthesis'),
    (r'climate|glaciology|regional geology|palaeoclimate|anthropogenic contaminants', 'cryosphere, climate and geology (context)'),
    (r'health|standards', 'health'),
    (r'remote sensing|permafrost biogeochemistry', 'remote sensing'),
    (r'palaeo|sediment', 'sediments and archives'),
    (r'geophysics|hydrology|hydrogeology|isotope', 'hydrology and geophysics'),
    (r'remediation|governance', 'management and remediation'),
    (r'.*', 'miscellaneous'),
]
def family(t):
    for f in ('Geomorphology', 'Hydrochemistry', 'Soil'):
        if t.startswith(f): return f
    return 'Other'
def harmonise(t):
    if family(t) != 'Other': return t
    lab = t.split('–', 1)[-1].strip()
    if lab == 'to be classified': return 'Other – to be classified'
    for pat, h in OTHER_MAP:
        if re.search(pat, lab, re.I): return 'Other – ' + h
changed = collections.Counter()
for x in REFS:
    x['theme_original'] = x['main_theme']
    x['family'] = family(x['main_theme'])
    x['theme'] = harmonise(x['main_theme'])
    if x['theme'] != x['main_theme']: changed[(x['main_theme'], x['theme'])] += 1
    x['themes_sec'] = list(dict.fromkeys(harmonise(t) for t in x['secondary_themes'] if harmonise(t) != x['theme']))
n_groups = len({b for _, b in changed} | {x['theme'] for x in REFS if x['family'] == 'Other'})
NOTES.append(f"The {len({a for a, _ in changed})} ad hoc “Other – …” main-theme labels were grouped into {n_groups} sub-themes "
             "(same grouping applied to secondary themes); the original label is kept in each note and in the table. "
             "Filters and colours use the 4 families Geomorphology / Hydrochemistry / Soil / Other.")

def country_list(p):
    return [s.strip() for s in re.split(r'[;,/]', p or '') if s.strip()]
for x in REFS:
    x['country_l'] = country_list(x['country'])

# formatting
def last(a): return a.split(',')[0].strip()
def short_auth(x):
    a = x['authors']
    if not a: return x['venue'] or 'Anonymous'
    if len(a) == 1: return last(a[0])
    if len(a) == 2: return f"{last(a[0])} & {last(a[1])}"
    return f"{last(a[0])} et al."
def label(x): return f"{short_auth(x)} {x['year']}"
def full_auth(x):
    a = x['authors']
    return '; '.join(a) if len(a) <= 12 else '; '.join(a[:10]) + f' … (+{len(a) - 10})'
def full_ref(x):
    s = f"{full_auth(x)} ({x['year']}). {x['title']}."
    v = x['venue']
    if v:
        s += f" {v}"
        if x['volume']: s += f", {x['volume']}"
        if x['issue']: s += f"({x['issue']})"
        if x['pages']: s += f", {x['pages']}"
        s += '.'
    if x['publisher'] and x['doc_type'] != 'article': s += f" {x['publisher']}."
    return s
def doi_url(x): return f"https://doi.org/{x['doi']}" if x['doi'] else (x['url'] or '')
for x in REFS:
    x['label'] = label(x); x['ref'] = full_ref(x)

# ---------- 2. reading table (CSV + XLSX) ----------
COLS = ['Key', 'Year', 'Authors', 'Title', 'Journal / venue', 'DOI', 'Document type', 'Main theme',
        'Secondary themes', 'Region', 'Country', 'Site', 'Origin', 'Trigger', 'Methods', 'Key results',
        'Reading level', 'Cites (in corpus)', 'Cited by (in corpus)', 'Status',
        'Main theme (original label)', 'Context', 'Peer reviewed', 'URL', 'Open access', 'Note']
def row(x):
    return [x['key'], x['year'], '; '.join(x['authors']), x['title'], x['venue'], x['doi'], x['doc_type'],
            x['theme'], '; '.join(x['themes_sec']), x['region'], x['country'], x['site'], x['origin'], x['trigger'],
            x['methods'], x['results'], x['reading_level'], '; '.join(x['cites']), '; '.join(x['cited_by']),
            x['status'], x['theme_original'], 'yes' if x['context'] else 'no', 'yes' if x['peer_reviewed'] else 'no',
            doi_url(x), x['oa_url'], x.get('note', '')]
ORDER = sorted(REFS, key=lambda x: (x['year'], x['key']))
with open(OUT / 'ARD_reading_table.csv', 'w', newline='', encoding='utf-8-sig') as f:
    w = csv.writer(f, delimiter=',', quoting=csv.QUOTE_MINIMAL)
    w.writerow(COLS)
    for x in ORDER: w.writerow(row(x))

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
wb = Workbook(); ws = wb.active; ws.title = 'Reading table'
F = 'Arial'
ws.append(COLS)
for x in ORDER: ws.append(row(x))
for c in ws[1]:
    c.font = Font(name=F, bold=True, color='FFFFFF', size=10); c.fill = PatternFill('solid', fgColor='2E3A36')
    c.alignment = Alignment(wrap_text=True, vertical='center')
widths = {'Key': 18, 'Year': 7, 'Authors': 30, 'Title': 50, 'Journal / venue': 26, 'DOI': 28, 'Document type': 14,
          'Main theme': 24, 'Secondary themes': 28, 'Region': 18, 'Country': 14, 'Site': 28, 'Origin': 10,
          'Trigger': 22, 'Methods': 50, 'Key results': 60, 'Reading level': 13, 'Cites (in corpus)': 28,
          'Cited by (in corpus)': 28, 'Status': 16, 'Main theme (original label)': 26, 'Context': 9,
          'Peer reviewed': 9, 'URL': 30, 'Open access': 30, 'Note': 50}
FAM_FILL = {'Geomorphology': 'DCE8F8', 'Hydrochemistry': 'FCE3D8', 'Soil': 'D6F2E7', 'Other': 'ECEDEB'}
for i, col in enumerate(COLS, 1):
    ws.column_dimensions[get_column_letter(i)].width = widths.get(col, 15)
thin = Side(style='thin', color='D0D4D0')
for r_i, x in enumerate(ORDER, 2):
    for c in ws[r_i]:
        c.font = Font(name=F, size=9); c.alignment = Alignment(wrap_text=True, vertical='top')
        c.border = Border(bottom=thin)
    ws.cell(r_i, COLS.index('Main theme') + 1).fill = PatternFill('solid', fgColor=FAM_FILL[x['family']])
    if x['doi']:
        c = ws.cell(r_i, COLS.index('DOI') + 1); c.hyperlink = doi_url(x); c.font = Font(name=F, size=9, color='1C5CAB', underline='single')
    if x['status'] == 'initial bibliography':
        ws.cell(r_i, COLS.index('Status') + 1).font = Font(name=F, size=9, bold=True)
    ws.row_dimensions[r_i].height = 60
ws.freeze_panes = 'E2'; ws.auto_filter.ref = ws.dimensions
ws.row_dimensions[1].height = 32
rd = wb.create_sheet('Read me')
lines = [('Reading table – ARD corpus (Zotero library)', True),
         (f"Generated from data/corpus.json (corpus dated {META['date']}): {len(REFS)} references, one row per document, sorted by year.", False),
         ('“Cites” and “Cited by” only cover references within the corpus (“Cited by” is the inverse of “Cites”).', False),
         ('“Reading level” states what was read (metadata, abstract, full text): the Methods and Key results columns say no more than that.', False),
         ('', False), ('Scope', True), (META['scope'], False), ('', False), ('Corpus choices', True)]
lines += [('• ' + c, False) for c in META['choices']]
lines += [('', False), ('Build choices', True)] + [('• ' + c, False) for c in NOTES]
lines += [('', False), ('Limitations', True)] + [('• ' + c, False) for c in META['limitations']]
for t, b in lines:
    rd.append([t]); c = rd.cell(rd.max_row, 1)
    c.font = Font(name=F, size=11 if b else 10, bold=b); c.alignment = Alignment(wrap_text=True, vertical='top')
rd.column_dimensions['A'].width = 130
wb.save(OUT / 'ARD_reading_table.xlsx')

# ---------- 3. BibLaTeX ----------
def bibesc(s):
    s = s or ''
    s = s.replace('\\', '\\textbackslash{}')
    for ch in '&%$#_':
        s = s.replace(ch, '\\' + ch)
    s = s.replace('{', '\\{').replace('}', '\\}').replace('\\textbackslash\\{\\}', '\\textbackslash{}')
    s = s.replace('~', '\\textasciitilde{}').replace('^', '\\textasciicircum{}')
    return s
TYPE = {'article': 'article', 'chapter': 'incollection', 'book': 'book', 'thesis': 'thesis', 'report': 'report',
        'conference abstract': 'inproceedings', 'preprint': 'online'}
def kw(x):
    k = ['theme: ' + x['theme']] + ['theme: ' + t for t in x['themes_sec']]
    k += ['location: ' + x['region']] + ['location: ' + p for p in x['country_l']]
    k += ['origin: ' + x['origin'], 'status: ' + x['status'], 'reading: ' + x['reading_level']]
    if x['context']: k.append('context')
    return ', '.join(dict.fromkeys(t.replace(',', ' ') for t in k))
bib = [f"% ARD_corpus.bib – {len(REFS)} references – generated from data/corpus.json (corpus dated {META['date']})",
       '% BibLaTeX format (UTF-8). Normalised keywords: "theme: …", "location: …", "origin: …", "status: …", "reading: …", "context".', '']
for x in ORDER:
    t = TYPE[x['doc_type']]
    f = [('author', ' and '.join(bibesc(a) for a in x['authors'])), ('title', '{' + bibesc(x['title']) + '}'),
         ('date', str(x['year']))]
    v = bibesc(x['venue']); pub = bibesc(x['publisher'])
    if t == 'article': f += [('journaltitle', v), ('volume', x['volume']), ('number', x['issue']), ('publisher', pub)]
    elif t == 'incollection': f += [('booktitle', v), ('publisher', pub)]
    elif t == 'book': f += [('series', v), ('publisher', pub)]
    elif t == 'thesis': f += [('institution', v or pub), ('type', 'Thesis')]
    elif t == 'report': f += [('type', v), ('institution', pub)]
    elif t == 'inproceedings': f += [('booktitle', v), ('publisher', pub)]
    elif t == 'online': f += [('organization', pub), ('pubstate', 'prepublished')]
    if x['pages']: f.append(('pages', bibesc(x['pages'].replace('–', '-').replace('--', '-').replace('-', '--'))))
    if x['doi']: f.append(('doi', x['doi']))
    u = x['oa_url'] or x['url']
    if u and not (x['doi'] and u == doi_url(x)): f.append(('url', bibesc(u)))
    elif not x['doi'] and x['url']: f.append(('url', bibesc(x['url'])))
    if t == 'online' and not any(k == 'url' for k, _ in f): f.append(('url', doi_url(x)))
    f.append(('keywords', bibesc(kw(x))))
    if x.get('note'): f.append(('annotation', bibesc(x['note'])))
    bib.append(f"@{t}{{{x['key']},")
    for k, val in f:
        if val: bib.append(f"  {k} = {{{val}}},")
    bib.append('}\n')
(OUT / 'ARD_corpus.bib').write_text('\n'.join(bib), encoding='utf-8')
NOTES.append("BibLaTeX types: article → @article, chapter → @incollection, book → @book, thesis → @thesis, report → @report, "
             "conference abstract → @inproceedings, preprint → @online (the types Better BibTeX exports). "
             "A “reading: …” keyword keeps the reading level visible in Zotero.")

# ---------- 4. page ----------
def pct(a, b): return f"{round(100 * a / b)}%" if b else '–'
rows = []
for rg in REGIONS:
    xs = [x for x in REFS if x['region'] == rg]
    if not xs: continue
    trig = collections.Counter(x['trigger'] if not x['trigger'].startswith('other') else 'other' for x in xs)
    trig.pop('not applicable', None)
    top = ', '.join(f"{k} ({n})" for k, n in trig.most_common(2)) or '–'
    ori = collections.Counter(x['origin'] for x in xs); ori.pop('not applicable', None)
    yrs = sorted(x['year'] for x in xs)
    rows.append(f"<tr><th scope='row'>{html.escape(rg)}</th><td class='n'>{len(xs)}</td><td class='n'>{sum(x['status'] == 'initial bibliography' for x in xs)}</td>"
                f"<td class='n'>{yrs[0]}–{yrs[-1]}</td><td class='n'>{pct(sum(x['reading_level'] == 'full text' for x in xs), len(xs))}</td>"
                f"<td>{html.escape(top)}</td><td>{html.escape(', '.join(f'{k} {n}' for k, n in ori.most_common()) or '–')}</td></tr>")
table = ("<div class='tablewrap'><table class='reg'><thead><tr><th>Region</th><th class='n'>Refs</th><th class='n'>Initial bibl.</th>"
         "<th class='n'>Years</th><th class='n'>Full text</th><th>Most cited triggers</th><th>Origin</th></tr></thead><tbody>"
         + ''.join(rows) + '</tbody></table></div>')
cross = sum(1 for x in REFS for k in x['cites'] if {x['region'], BY[k]['region']} == {'Pyrenees', 'Cordillera Blanca'})
syn = (DATA / 'synthesis.html').read_text(encoding='utf-8').replace('<!--TABLE_REGIONS-->', table).replace('<!--CROSS_PYR_CB-->', str(cross))
missing = set()
def link(m):
    k = m.group(1)
    if k not in BY: missing.add(k); return m.group(0)
    x = BY[k]
    return f"<a class='rk f-{x['family']}' href='#{k}' data-k='{k}'>{html.escape(x['label'])}</a>"
syn = re.sub(r'\[\[([A-Za-z0-9]+)\]\]', link, syn)
if missing:
    sys.exit(f'Unknown keys in data/synthesis.html: {sorted(missing)}')
SYN_KEYS = set(re.findall(r"data-k='([A-Za-z0-9]+)'", syn))

data = [{
    'k': x['key'], 'y': x['year'], 'l': x['label'], 'ref': x['ref'], 't': x['title'], 'au': x['authors'],
    'doi': x['doi'], 'u': doi_url(x), 'oa': x['oa_url'], 'v': x['venue'], 'dt': x['doc_type'], 'pr': x['peer_reviewed'],
    'st': x['status'], 'ctx': x['context'], 'f': x['family'], 'th': x['theme'], 'tho': x['theme_original'], 'ts': x['themes_sec'],
    'rg': x['region'], 'py': x['country_l'], 'pys': x['country'], 'si': x['site'], 'or': x['origin'], 'de': x['trigger'],
    'q': x['question'], 'm': x['methods'], 'r': x['results'], 'a': x['contribution'], 'nl': x['reading_level'],
    'c': x['cites'], 'cb': x['cited_by'], 'rq': x.get('note', ''), 'fv': x['found_via']} for x in ORDER]
meta = {'date': META['date'], 'scope': META['scope'], 'choices': META['choices'], 'limitations': META['limitations'],
        'notes2': NOTES, 'sources': META['sources_used'], 'nq': len(META.get('queries', [])), 'regions': REGIONS}
tpl = (ROOT / 'templates' / 'map_template.html').read_text(encoding='utf-8')
page = (tpl.replace('/*__DATA__*/null', json.dumps(data, ensure_ascii=False).replace('</', '<\\/'))
           .replace('/*__META__*/null', json.dumps(meta, ensure_ascii=False).replace('</', '<\\/'))
           .replace('<!--__SYNTHESIS__-->', syn))
(FRAG / 'ard_map_fragment.html').write_text(page, encoding='utf-8')
# stand-alone page for GitHub Pages: title, styles and links go into <head>
cut = page.index('<div id="tip"')
(OUT / 'index.html').write_text('<!doctype html>\n<html lang="en">\n<head>\n<meta charset="utf-8">\n'
    '<meta name="viewport" content="width=device-width,initial-scale=1,viewport-fit=cover">\n'
    '<style>:root{padding-top:env(safe-area-inset-top,0px);padding-bottom:env(safe-area-inset-bottom,0px)}'
    'img{max-width:100%}[hidden]{display:none!important}</style>\n'
    + page[:cut] + '\n</head>\n<body>\n' + page[cut:] + '\n</body>\n</html>\n', encoding='utf-8')
(OUT / '.nojekyll').write_text('')

# ---------- 5. candidates page (latest literature watch) ----------
import os
def wrap_page(fragment, marker):
    cut = fragment.index(marker)
    return ('<!doctype html>\n<html lang="en">\n<head>\n<meta charset="utf-8">\n'
            '<meta name="viewport" content="width=device-width,initial-scale=1,viewport-fit=cover">\n'
            '<style>:root{padding-top:env(safe-area-inset-top,0px);padding-bottom:env(safe-area-inset-bottom,0px)}'
            'img{max-width:100%}[hidden]{display:none!important}</style>\n'
            + fragment[:cut] + '\n</head>\n<body>\n' + fragment[cut:] + '\n</body>\n</html>\n')

CAND = DATA / 'candidates'
runs = sorted(p for p in CAND.glob('*.json') if re.fullmatch(r'\d{4}-\d{2}-\d{2}', p.stem))
run = json.loads(runs[-1].read_text(encoding='utf-8')) if runs else {}
pending_ids, decided_ids = set(), set()
if (CAND / 'to_screen.csv').exists():
    with open(CAND / 'to_screen.csv', encoding='utf-8-sig', newline='') as f:
        for l in csv.DictReader(f):
            (decided_ids if (l.get('decision') or '').strip() else pending_ids).add(l.get('openalex', ''))
rejected = set()
if (CAND / 'rejected.txt').exists():
    rejected = {l.strip().lower() for l in (CAND / 'rejected.txt').read_text(encoding='utf-8').splitlines() if l.strip()}
corpus_oa = {x['openalex'] for x in REFS if x.get('openalex')}
corpus_dois = {x['doi'].lower() for x in REFS if x.get('doi')}
cands = []
for c in run.get('candidates', []):
    oid, doi = c.get('openalex', ''), (c.get('doi') or '').lower()
    if oid in corpus_oa or (doi and doi in corpus_dois): st = 'added'
    elif (doi and doi in rejected) or oid.rsplit('/', 1)[-1].lower() in rejected: st = 'rejected'
    elif oid in decided_ids: st = 'decided'
    else: st = 'pending'
    cands.append({'s': c['score'], 'y': c.get('year'), 't': c.get('title', ''), 'au': c.get('authors', ''), 'v': c.get('venue', ''),
                  'doi': c.get('doi', ''), 'oa': oid, 'oa_url': c.get('open_access', ''), 'key': c.get('proposed_key', ''),
                  'via': [v for v in (c.get('found_via') or '').split(' ; ') if v],
                  'cites': [k for k in (c.get('cites_in_corpus') or '').split('; ') if k], 'st': st})
cmeta = {'date': run.get('date', ''), 'since': run.get('since', ''), 'errors': run.get('errors', []),
         'repo': os.environ.get('GITHUB_REPOSITORY', ''),
         'labels': {x['key']: x['label'] for x in REFS},
         'themes': MAIN_THEMES + ['Other – to be classified'], 'regions': REGIONS[1:]}
ctpl = (ROOT / 'templates' / 'candidates_template.html').read_text(encoding='utf-8')
cpage = (ctpl.replace('/*__CANDS__*/null', json.dumps(cands, ensure_ascii=False).replace('</', '<\\/'))
             .replace('/*__META__*/null', json.dumps(cmeta, ensure_ascii=False).replace('</', '<\\/')))
(FRAG / 'candidates_fragment.html').write_text(cpage.replace('<!--BODY-->', ''), encoding='utf-8')
(OUT / 'candidates.html').write_text(wrap_page(cpage, '<!--BODY-->').replace('<!--BODY-->', ''), encoding='utf-8')
print(f"Candidates page: {len(cands)} candidates from the watch of {run.get('date') or '(none yet)'}, "
      f"{sum(c['st'] == 'pending' for c in cands)} waiting for a decision.")

# ---------- 6. consistency check ----------
import openpyxl
xl = openpyxl.load_workbook(OUT / 'ARD_reading_table.xlsx')['Reading table']
xl_keys = {r[0] for r in xl.iter_rows(min_row=2, values_only=True)}
xl_dois = {(r[5] or '').lower() for r in xl.iter_rows(min_row=2, values_only=True)} - {''}
csv_rows = list(csv.DictReader(open(OUT / 'ARD_reading_table.csv', encoding='utf-8-sig')))
csv_keys = {r['Key'] for r in csv_rows}; csv_dois = {r['DOI'].lower() for r in csv_rows} - {''}
btxt = (OUT / 'ARD_corpus.bib').read_text(encoding='utf-8')
bib_keys = set(re.findall(r'^@\w+\{([^,]+),', btxt, re.M)); bib_dois = {d.lower() for d in re.findall(r'^  doi = \{([^}]+)\}', btxt, re.M)}
pg = json.loads(re.search(r'const DATA = (\[.*?\]);\n', page, re.S).group(1))
pg_keys = {d['k'] for d in pg}; pg_dois = {d['doi'].lower() for d in pg} - {''}
ref_dois = {x['doi'].lower() for x in REFS} - {''}
res = {n: (k == KEYS, d == ref_dois, len(k), len(d)) for n, k, d in
       [('xlsx', xl_keys, xl_dois), ('csv', csv_keys, csv_dois), ('bib', bib_keys, bib_dois), ('page', pg_keys, pg_dois)]}
print(json.dumps(res))
if not all(a and b for a, b, _, _ in res.values()):
    sys.exit(f'Deliverables do not match: {res}')
json.dump({'n': len(REFS), 'notes': NOTES, 'cross': cross, 'syn_keys': len(SYN_KEYS)},
          open(FRAG / 'build_info.json', 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
print(f'OK: {len(REFS)} references; {len(SYN_KEYS)} cited in the synthesis; Pyrenees–Cordillera Blanca citation links: {cross}')
